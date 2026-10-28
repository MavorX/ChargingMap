"""Regression checks for the actual company template, not just its colors."""

import contextlib
import hashlib
import importlib.util
import io
import os
from pathlib import Path
import unittest
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches


SPEC = importlib.util.spec_from_file_location("presentation_builder", Path(__file__).with_name("build-presentation.py"))
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)
TEMPLATE = Path(os.environ.get("CHARGINGMAP_TEMPLATE", str(BUILDER.TEMPLATE)))
DECK = BUILDER.OUT / "ChargingMap_Project_Overview_ZH-EN.pptx"


class PresentationTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not TEMPLATE.is_file():
            raise unittest.SkipTest("The company template is local/private; provide CHARGINGMAP_TEMPLATE.")

    def setUp(self):
        self.prs = Presentation(DECK)

    def validate(self):
        with contextlib.redirect_stdout(io.StringIO()):
            BUILDER.validate_branding(self.prs, TEMPLATE)

    def master_logo(self, name):
        return next(shape for shape in self.prs.slide_masters[0].shapes if shape.name == name)

    def test_original_layouts_and_both_logos_are_kept(self):
        self.validate()
        with ZipFile(DECK) as deck, ZipFile(TEMPLATE) as template:
            self.assertEqual(deck.read("ppt/theme/theme1.xml"), template.read("ppt/theme/theme1.xml"))

    def test_missing_company_wordmark_is_rejected(self):
        logo = self.master_logo("Picture 30")
        logo._element.getparent().remove(logo._element)
        with self.assertRaisesRegex(AssertionError, "Missing company wordmark or RISE UP"):
            self.validate()

    def test_missing_rise_up_is_rejected(self):
        logo = self.master_logo("Grafik 25")
        logo._element.getparent().remove(logo._element)
        with self.assertRaisesRegex(AssertionError, "Missing company wordmark or RISE UP"):
            self.validate()

    def test_moved_logo_is_rejected(self):
        self.master_logo("Picture 30").left += Inches(0.1)
        with self.assertRaisesRegex(AssertionError, "Altered template logo"):
            self.validate()

    def test_logo_overlay_is_rejected(self):
        self.prs.slides[1].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.4), Inches(0.1), Inches(2.4), Inches(0.6))
        with self.assertRaisesRegex(AssertionError, "covers the original"):
            self.validate()

    def test_altered_logo_asset_is_rejected(self):
        logo = next(shape for shape in self.prs.slides[0].slide_layout.shapes if shape.name == "Picture 11")
        embedded = next(element.get(qn("r:embed")) for element in logo._element.iter() if element.get(qn("r:embed")))
        image = logo.part.related_part(embedded)
        image._blob = image.blob + b"modified-in-memory-for-regression"
        with self.assertRaisesRegex(AssertionError, "Altered template logo"):
            self.validate()

    def test_no_manual_or_ole_is_distributed(self):
        with ZipFile(DECK) as archive:
            names = archive.namelist()
            self.assertIn(b"<Slides>10</Slides>", archive.read("docProps/app.xml"))
            self.assertFalse(any("/embeddings/" in name or "/tags/" in name or name.startswith("customXml/") for name in names))
            self.assertEqual(len([name for name in names if name.startswith("ppt/slideLayouts/slideLayout") and name.endswith(".xml")]), 4)
            for name in names:
                if name.endswith(".xml"):
                    content = archive.read(name)
                    for forbidden in [b"User Manual", b"Click to edit", b"Department K-OE", b"firstname.lastname@emailaddress.com"]:
                        self.assertNotIn(forbidden, content, name)

    def test_template_source_is_not_modified(self):
        before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        loaded = BUILDER.load_template(TEMPLATE)
        self.assertEqual(len(loaded.slides), 0)
        self.assertEqual(len(loaded.slide_layouts), 4)
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), before)
        with self.assertRaisesRegex(FileNotFoundError, "original company template is required"):
            BUILDER.load_template(TEMPLATE.with_name("missing-corporate-template.potx"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
