"""Build the editable bilingual ChargingMap project deck.

Requires the user-provided VCTC template, python-pptx and Pillow. Optional
PDF/preview rendering also uses LibreOffice and PyMuPDF. The original template
and font files are not redistributed; its used masters/logos remain in the deck.
"""

import argparse
from io import BytesIO
import hashlib
import json
import math
import os
from pathlib import Path
import re
import subprocess
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image, ImageDraw
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "presentation"
ASSETS = OUT / "assets"
TEMPLATE = OUT / "VCTC Template_Rise Up.potx"
W, H = 13.333333, 7.5
INK, GREEN, NEON = "002733", "008C82", "C2FE06"
MINT, GRAY, LINE, BG, WHITE = "99D1CD", "4C6870", "CCD3D6", "FFFFFF", "FFFFFF"
CN, EN, HEAD = "Microsoft YaHei", "+mn-lt", "+mj-lt"
DEMO = "https://mavorx.github.io/ChargingMap/"
FOOTER_TYPES = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}
LAYOUT_PARTS = {"cover": "slideLayout3.xml", "light": "slideLayout24.xml",
                "dark": "slideLayout23.xml", "final": "slideLayout40.xml"}


def template_presentation(source):
    """Open POTX in memory, retaining real layout/master artwork and geometry."""
    if not source.is_file():
        raise FileNotFoundError(f"The original company template is required: {source}")
    converted = BytesIO()
    with ZipFile(source) as original, ZipFile(converted, "w", ZIP_DEFLATED) as output:
        for part in original.infolist():
            payload = original.read(part.filename)
            if part.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                )
            output.writestr(part, payload)
    converted.seek(0)
    return Presentation(converted)


def prune_unreferenced_relationships(part):
    """Remove unused add-in, OLE and guide assets, never inherited logo images."""
    for rel in list(part.rels.values()):
        if not rel.reltype.endswith(("/image", "/oleObject", "/tags", "/customXml")):
            continue
        if not any(value == rel.rId for element in part._element.iter()
                   for name, value in element.attrib.items()
                   if name.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")):
            part.drop_rel(rel.rId)


def load_template(source):
    prs = template_presentation(source)
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    layouts = {role: next(layout for layout in prs.slide_layouts
                          if Path(str(layout.part.partname)).name == name)
               for role, name in LAYOUT_PARTS.items()}
    # Original sample slides and user-manual layouts are not project content.
    for item in list(prs.slides._sldIdLst):
        prs.part.drop_rel(item.rId)
        prs.slides._sldIdLst.remove(item)
    for element in list(prs.part._element):
        if element.tag in {qn("p:extLst"), qn("p:sectionLst")}:
            prs.part._element.remove(element)
    master = prs.slide_masters[0]
    keep = {layout.part for layout in layouts.values()}
    for item in list(master.slide_layouts._sldLayoutIdLst):
        if master.part.related_part(item.rId) not in keep:
            master.part.drop_rel(item.rId)
            master.slide_layouts._sldLayoutIdLst.remove(item)
    for shape in list(master.shapes):
        if not shape.is_placeholder and shape.name not in {"Picture 30", "Grafik 25"}:
            # Only off-canvas guide marks, old classification and think-cell data.
            master.shapes._spTree.remove(shape._element)
    for container in [master, *layouts.values()]:
        for shape in container.shapes:
            if shape.is_placeholder and shape.has_text_frame:
                shape.text = ""  # Remove instructions but retain placeholder styles.
        prune_unreferenced_relationships(container.part)
    for shape in prs.notes_master.shapes:
        if shape.has_text_frame:
            shape.text = ""  # Keep notes geometry/styles, not the template's sample text/date.
    for part in list(prs.part.package.iter_parts()):
        for rel in list(part.rels.values()):
            if rel.reltype.endswith(("/customXml", "/tags", "/oleObject", "/thumbnail")):
                if hasattr(part, "_element"):
                    for element in list(part._element.iter()):
                        if rel.rId in [value for name, value in element.attrib.items()
                                       if name.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")]:
                            element.getparent().remove(element)
                part.drop_rel(rel.rId)
    for rel in list(prs.part.package._rels.values()):
        if rel.reltype.endswith(("/custom-properties", "/thumbnail")):
            prs.part.package.drop_rel(rel.rId)
    prs._charging_layouts = layouts
    return prs


def logo_signature(shape):
    assets = []
    for element in shape._element.iter():
        relationship = element.get(qn("r:embed"))
        if relationship:
            payload = shape.part.related_part(relationship).blob
            assets.append(hashlib.sha256(payload).hexdigest())
    return (shape.name, shape.left, shape.top, shape.width, shape.height, tuple(assets))


def logo_shapes(slide):
    layout = slide.slide_layout
    shapes = list(layout.shapes)
    if layout._element.get("showMasterSp") != "0":
        shapes += list(layout.slide_master.shapes)
    return [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]


def validate_branding(prs, source):
    """Gate every slide on original template assets/geometry and clear areas."""
    reference = template_presentation(source)
    originals = {str(layout.part.partname): layout for layout in reference.slide_layouts}
    assert len(prs.slide_masters) == 1 and len(prs.slide_layouts) == 4
    assert len(prs.slides) == 10
    checked = 0
    for index, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout
        role = "cover" if index == 1 else "final" if index == 10 else "dark" if index == 3 else "light"
        assert Path(str(layout.part.partname)).name == LAYOUT_PARTS[role], f"Wrong template layout: {index}"
        original = originals[str(layout.part.partname)]
        actual_bg, original_bg = layout._element.cSld.bg, original._element.cSld.bg
        assert (actual_bg.xml if actual_bg is not None else None) == (original_bg.xml if original_bg is not None else None)
        expected = list(original.shapes)
        if original._element.get("showMasterSp") != "0":
            expected += list(original.slide_master.shapes)
        expected = [shape for shape in expected if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        marks = logo_shapes(slide)
        assert len(marks) == len(expected) == 2, f"Missing company wordmark or RISE UP: slide {index}"
        assert sorted(map(logo_signature, marks)) == sorted(map(logo_signature, expected)), f"Altered template logo: slide {index}"
        artwork = [etree.tostring(shape._element, method="c14n") for shape in layout.shapes if not shape.is_placeholder]
        original_artwork = [etree.tostring(shape._element, method="c14n") for shape in original.shapes if not shape.is_placeholder]
        assert artwork == original_artwork, f"Original layout artwork was changed: slide {index}"
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type in FOOTER_TYPES:
                continue  # Original right-aligned footer boxes extend behind the logo region.
            for mark in marks:
                overlaps = (shape.left < mark.left + mark.width and shape.left + shape.width > mark.left
                            and shape.top < mark.top + mark.height and shape.top + shape.height > mark.top)
                assert not overlaps, f"Slide {index}: {shape.name} covers the original {mark.name}"
        checked += len(marks)
    print(f"Template check passed: {checked} original logo instances on 10 slides; 4 genuine layouts; no logo overlap.")


def update_document_properties(prs):
    # python-pptx preserves this generic part verbatim; the source still lists
    # all 73 sample slide titles unless explicitly replaced after composition.
    namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    root = etree.Element(f"{{{namespace}}}Properties", nsmap={None: namespace})
    values = {"Application": "ChargingMap Presentation Builder", "PresentationFormat": "Widescreen",
              "Slides": str(len(prs.slides)), "Notes": str(len(prs.slides)), "HiddenSlides": "0"}
    for name, value in values.items():
        etree.SubElement(root, f"{{{namespace}}}{name}").text = value
    for rel in prs.part.package._rels.values():
        if rel.reltype.endswith("/extended-properties"):
            rel.target_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill, stroke=None):
    # The supplied template explicitly specifies sharp-corner text/image boxes.
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    effect_ref = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style/{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
    if effect_ref is not None:
        effect_ref.set("idx", "0")
    return shape


def text(slide, value, x, y, w, h, size=18, fill=INK, bold=False,
         font=CN, align=PP_ALIGN.LEFT, link=None, shape=None):
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.12
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(fill)
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", CN)
        run._r.get_or_add_rPr().append(east_asian)
    if link:
        # A shape-level action keeps the intended brand color instead of the
        # viewer's default blue hyperlink run style.
        shape.click_action.hyperlink.address = link
    return shape


def pair(slide, zh, en, x, y, w, size=18, fill=INK, secondary=GRAY, center=False):
    align = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    text(slide, zh, x, y, w, 0.4, size=size, fill=fill, bold=True, align=align)
    text(slide, en, x, y + 0.43, w, 0.49, size=12.5, fill=secondary, font=EN, align=align)


def pill(slide, value, x, y, w, fill=GREEN, fg=WHITE):
    box(slide, x, y, w, 0.35, fill)
    text(slide, value, x + 0.1, y + 0.055, w - 0.2, 0.24, size=10, fill=fg, font=EN,
         align=PP_ALIGN.CENTER)


def photo(slide, name, x, y, w, crop=None, dark=False):
    source = ASSETS / name
    with Image.open(source) as image:
        iw, ih = image.size
        if crop:
            image = image.crop(crop)
            cached = ASSETS / (source.stem + "-crop.png")
            image.save(cached, optimize=True)
            source = cached
            iw, ih = image.size
    h = w * ih / iw
    box(slide, x - 0.018, y - 0.018, w + 0.036, h + 0.036, INK if dark else LINE)
    slide.shapes.add_picture(str(source), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return h


def phone(slide, name, x, y, w=2.0):
    return photo(slide, name, x, y, w)


def new_slide(prs, number, zh, en, section, dark=False):
    role = "cover" if number == 1 else "final" if number == 10 else "dark" if dark else "light"
    layout = prs._charging_layouts[role]
    slide = prs.slides.add_slide(layout)
    # Do not paint a new slide background or copy over the original brand areas.
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            slide.shapes._spTree.remove(shape._element)
    fg, muted = (WHITE, MINT) if dark else (INK, GRAY)
    if number not in {1, 10}:
        heading = next(p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.TITLE)
        action = next(p for p in layout.placeholders if p.name == "Action Title")
        slide.shapes.clone_placeholder(heading)
        title_shape = slide.placeholders[heading.placeholder_format.idx]
        title_shape.width = Inches(9.7)  # Keep the top-right wordmark clear.
        text(slide, zh, 0, 0, 0, 0, size=26, fill=fg, bold=True, shape=title_shape)
        slide.shapes.clone_placeholder(action)
        subtitle = slide.placeholders[action.placeholder_format.idx]
        text(slide, en, 0, 0, 0, 0, size=15, fill=muted if dark else GREEN,
             font=EN, shape=subtitle)
    for placeholder in layout.placeholders:
        kind = placeholder.placeholder_format.type
        if kind not in FOOTER_TYPES:
            continue
        slide.shapes.clone_placeholder(placeholder)
        target = slide.placeholders[placeholder.placeholder_format.idx]
        value = (str(number) if kind == PP_PLACEHOLDER.SLIDE_NUMBER else
                 "2026.10.28" if kind == PP_PLACEHOLDER.DATE else
                 "ChargingMap  |  Project Overview / 项目介绍  |  DEMO")
        text(slide, value, 0, 0, 0, 0, size=7.5, fill="99A9AD" if dark else GRAY,
             font=EN, align=PP_ALIGN.RIGHT, shape=target)
    return slide


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def plot_distribution():
    """Plot committed coordinates and approximate road corridors, not live POIs."""
    stations = json.loads((ROOT / "public/data/station_list.json").read_text(encoding="utf-8"))
    stations = [s for s in stations if s["stationId"].startswith("SH")]
    assert len(stations) == 1500
    roads_source = (ROOT / "scripts/shanghai-roads.mjs").read_text(encoding="utf-8")
    roads = [tuple(float(v) for v in match) for match in re.findall(
        r"start:\s*\[([\d.]+),\s*([\d.]+)\],\s*end:\s*\[([\d.]+),\s*([\d.]+)\]", roads_source)]
    image = Image.new("RGB", (1680, 1040), "#" + INK)
    draw = ImageDraw.Draw(image)

    def panel(bounds, region, dot_size):
        x, y, width, height = region
        west, south, east, north = bounds
        scale = min(width / ((east - west) * math.cos(math.radians(31.2))), height / (north - south))
        px_width = (east - west) * math.cos(math.radians(31.2)) * scale
        px_height = (north - south) * scale
        ox, oy = x + (width - px_width) / 2, y + (height - px_height) / 2

        def xy(lng, lat):
            return ox + (lng - west) * math.cos(math.radians(31.2)) * scale, oy + (north - lat) * scale

        for lng1, lat1, lng2, lat2 in roads:
            if all(west <= lng <= east and south <= lat <= north for lng, lat in [(lng1, lat1), (lng2, lat2)]):
                draw.line([xy(lng1, lat1), xy(lng2, lat2)], fill="#33545F", width=2)
        for station in stations:
            lng, lat = station["lng"], station["lat"]
            if west <= lng <= east and south <= lat <= north:
                px, py = xy(lng, lat)
                draw.ellipse((px - dot_size, py - dot_size, px + dot_size, py + dot_size), fill="#" + NEON)

    panel((121.05, 30.68, 121.99, 31.88), (65, 35, 900, 955), 3.3)
    draw.rounded_rectangle((1005, 150, 1625, 865), radius=20, fill="#0C3440", outline="#33545F", width=2)
    panel((121.35, 31.15, 121.62, 31.34), (1030, 195, 570, 620), 3.5)
    image.save(ASSETS / "synthetic-distribution.png", optimize=True)


def build(template=TEMPLATE):
    ASSETS.mkdir(parents=True, exist_ok=True)
    required = ["desktop-light.png", "desktop-dark.png", "desktop-filter.png", "mobile-detail.png",
                "mobile-charging.png", "mobile-plug.png", "mobile-unplug.png", "mobile-paid.png",
                "mobile-english.png", "mobile-charging-dark.png", "english-detail-1440.png"]
    missing = [name for name in required if not (ASSETS / name).is_file()]
    if missing:
        raise FileNotFoundError(f"Capture/copy presentation screenshots first: {missing}")
    plot_distribution()
    prs = load_template(template)
    prs.core_properties.title = "VWer-Charging | 项目介绍 / Project Overview"
    prs.core_properties.subject = "Designed synthetic data, dual-theme maps, and a connected charging demo"
    prs.core_properties.author = "ChargingMap"
    prs.core_properties.last_modified_by = "ChargingMap"
    prs.core_properties.keywords = "ChargingMap, EV, Vue, AMap, bilingual, synthetic data"
    prs.core_properties.comments = "Editable bilingual presentation. All stations and transactions are simulated."

    # 01 — Hero
    s = new_slide(prs, 1, "", "", "PROJECT OVERVIEW", dark=True)
    text(s, "VWer-Charging", 0.446, 1.86, 6.4, 0.83, size=40, fill=WHITE, font=HEAD)
    text(s, "让每一次出发，都满电。", 0.446, 2.98, 6.25, 0.6, size=26, fill=WHITE)
    text(s, "One map.\nA connected charging journey.", 0.446, 3.82, 6.1, 0.96, size=23, fill=MINT, font=EN)
    photo(s, "desktop-light.png", 7.08, 2.0, 5.65, dark=True)
    text(s, "2026.10 / 项目介绍 / Project Overview", 0.446, 5.31, 6.2, 0.3, size=12.5, fill=MINT)
    text(s, "模拟站点与交易 / Simulated stations & transactions", 7.1, 6.12, 5.65, 0.26,
         size=10.5, fill=MINT, font=EN)
    notes(s, "项目定位：Vue + Pinia + 高德地图的充电服务演示，不控制真实设备、不发生真实扣款。\n"
          "A front-end EV charging demo, not a live operator platform. Screenshots were captured locally in Edge with the real AMap SDK.\n"
          "本稿直接使用用户提供的 VCTC Template_Rise Up.potx，保留原始母版、封面曲线、右上角公司字标和左下角 RISE UP 标识；不分发模板原件、手册页或字体文件。")

    # 02 — The journey
    s = new_slide(prs, 2, "从找站，到完成充电", "One connected journey — from discovery to completion.", "EXPERIENCE")
    for i, (zh, en) in enumerate([( "找得到", "Discover the right station"), ("接得上", "Select, plug in, then start"), ("结得清", "Stop, unplug, pay & review")]):
        y = 2.14 + i * 1.36
        text(s, f"0{i + 1}", 0.6, y, 0.62, 0.48, size=24, fill=GREEN, font=EN, bold=True)
        pair(s, zh, en, 1.4, y, 3.25, size=23)
    for name, x, zh, en in [("mobile-detail.png", 5.0, "发现站点", "DISCOVER"),
                            ("mobile-charging.png", 7.62, "充电进度", "CHARGE"),
                            ("mobile-paid.png", 10.24, "完成支付", "COMPLETE")]:
        phone(s, name, x, 1.98, 2.03)
        text(s, f"{zh}  /  {en}", x - 0.12, 6.55, 2.35, 0.3, size=11.5, fill=GRAY, font=CN, align=PP_ALIGN.CENTER)
    notes(s, "贯穿找站、组合筛选、详情、预约、明确选桩、接枪检测、充电、停止、拔枪、账单、支付与自愿评价。\n"
          "用户返回地图后仍能恢复当前会话；刷新恢复已保存状态并提示中断，不补算离线充电。\n"
          "End-to-end browser tests exercise the real UI controls. Payments, rewards and records remain local demo data.")

    # 03 — Spatial simulation
    s = new_slide(prs, 3, "模拟数据，不是随机撒点", "Designed synthetic data — not uniform random dots.", "SYNTHETIC DATA", dark=True)
    photo(s, "synthetic-distribution.png", 0.55, 1.98, 7.32, dark=True)
    text(s, "1,500", 0.8, 2.18, 2.25, 0.57, size=34, fill=NEON, font=EN, bold=True)
    text(s, "上海模拟站点 / Shanghai demo sites", 0.82, 2.86, 4.6, 0.3, size=12, fill=MINT)
    text(s, "中心城区局部\nCentral-area detail", 5.15, 2.14, 2.32, 0.66, size=11, fill=MINT, align=PP_ALIGN.CENTER)
    for i, (zh, en) in enumerate([("道路骨架", "Road-guided sampling"), ("多中心需求", "Polycentric demand"), ("街区疏密变化", "Neighborhood variation"), ("停车区域偏移", "Parking-area offsets")]):
        y = 2.1 + i * 1.08
        text(s, f"0{i + 1}", 8.26, y + 0.03, 0.55, 0.38, size=18, fill=NEON, font=EN, bold=True)
        pair(s, zh, en, 9.03, y, 3.55, size=20, fill=WHITE, secondary=MINT)
    text(s, "110 m 最小间距 / min gap  ·  ≤ 5 站 / sites within 250 m  ·  35–180 m 偏移 / offset", 0.62, 6.64, 12.0, 0.22, size=11.5, fill=MINT)
    notes(s, "图中直接绘制仓库内上海 1500 个模拟坐标及近似道路，右侧为中心城区局部；不是测绘数据或真实 POI。\n"
          "生成器按道路分段长度、道路权重、多中心需求高斯叠加与连续街区噪声抽样，并偏移到道路两侧 35–180 米的假设停车区域。\n"
          "110 米只防近乎重叠，不是目标间距。任一站点 250 米内最多 5 站（含自身），不为凑数量放宽约束。\n"
          "Sources: scripts/shanghai-layout.mjs, scripts/shanghai-demand.mjs, scripts/shanghai-roads.mjs; data-generation.spec.js.")

    # 04 — Data pipeline
    s = new_slide(prs, 4, "可复现的数据，按需抵达", "Reproducible data. On-demand detail.", "DATA ENGINEERING")
    cards = [("42", "固定随机种子", "Deterministic seed"), ("2,000", "上海 + 合肥", "1,500 Shanghai + 500 Hefei"),
             ("20", "每片 100 站", "100 stations per detail shard"), ("1", "同片请求共享", "One in-flight request per shard")]
    for i, (metric, zh, en) in enumerate(cards):
        x = 0.55 + 3.14 * i
        box(s, x, 2.22, 2.8, 2.86, WHITE, LINE)
        text(s, metric, x + 0.22, 2.53, 2.4, 0.94, size=48, fill=GREEN, font=EN, bold=True)
        pair(s, zh, en, x + 0.23, 3.69, 2.34, size=18)
        if i < 3:
            text(s, "→", x + 2.86, 3.28, 0.28, 0.4, size=18, fill=GRAY, font=EN)
    box(s, 0.55, 5.67, 12.22, 0.93, INK)
    text(s, "同一参数，逐字节一致", 0.85, 5.9, 5.8, 0.39, size=21, fill=WHITE, bold=True)
    text(s, "Same parameters. Byte-identical output.", 6.55, 5.96, 5.8, 0.34, size=17, fill=MINT, font=EN)
    notes(s, "默认 seed=42，上海 1500 站、合肥 500 站；两个地区使用独立的随机流。相同参数重复运行逐字节一致，上海数量变化不影响合肥数据。\n"
          "轻量列表先加载，详情按 stationId 推导到 [sh-]part-NN.json；每片 100 站，Promise 级缓存合并并发请求，失败可重试。\n"
          "The '1' describes a shared concurrent request per detail shard, not one network request for the whole app.\n"
          "Sources: generate_stations.js, src/services/api.js, data-generation.spec.js, api-details-cache.spec.js.")

    # 05 — Dual-theme map
    s = new_slide(prs, 5, "双色地图，只换观感", "Two themes. One map. No style reload.", "DUAL-THEME MAP")
    pill(s, "白天 / LIGHT", 0.55, 1.9, 1.58)
    pill(s, "夜间 / DARK", 6.87, 1.9, 1.58, fill=INK)
    photo(s, "desktop-light.png", 0.55, 2.4, 5.9)
    photo(s, "desktop-dark.png", 6.87, 2.4, 5.9)
    text(s, "0", 0.6, 6.55, 0.5, 0.33, size=22, fill=GREEN, font=EN, bold=True)
    text(s, "新增样式 / 瓦片请求  ·  New style / tile requests during toggles", 1.21, 6.58, 11.5, 0.26, size=12, fill=GRAY)
    notes(s, "技术方案：固定复用 amap://styles/dark 和同一个地图实例，白天仅对底图 canvas 做局部 CSS 明暗合成。不是逐次调用 setMapStyle，也没有另一张隐藏地图。\n"
          "DOM 标记、信息窗及版权区域不随底图反色。相机、标记节点和测试折线保持不变。\n"
          "本轮 Edge 实测：1440×1000 与 390×844 各连续切换 8 次，并阻断远程样式/瓦片服务后再切换；新增相应请求为 0。\n"
          "This is a local compositing approach, not two independent AMap styles. Two-frame observations are environment-dependent, not a frame-rate guarantee. Test polyline retention does not revalidate the live driving API.\n"
          "Sources: src/components/map/MapView.vue, map-theme.spec.js, scripts/capture-presentation.mjs, TEST_REPORT.md.")

    # 06 — Discovery synchronization
    s = new_slide(prs, 6, "地图、列表、筛选，一个答案", "Map, list and filters stay in sync.", "CONNECTED DISCOVERY")
    photo(s, "desktop-light.png", 0.55, 1.98, 7.02)
    photo(s, "desktop-filter.png", 7.95, 1.98, 2.07, crop=(1020, 64, 1440, 1000))
    for i, (zh, en) in enumerate([("视野联动", "Viewport-aware results"), ("组合筛选", "Combine availability,\nconnector & power"), ("可用数同步", "Live marker &\ncluster updates")]):
        y = 2.25 + i * 1.36
        box(s, 10.55, y + 0.07, 0.07, 0.32, GREEN)
        pair(s, zh, en, 10.8, y, 2.05, size=19)
    notes(s, "搜索针对全部站点的站名/地址，支持中英关键词；再按地图视野与分页展示。所有入口共享同一搜索与筛选状态。\n"
          "150ms 防抖合并视野更新；可用枪数变化也会同步筛选成员、标记与聚合总数。\n"
          "本轮修复了 AMap 重用标记时点击回调叠加、旧站点绑定，以及可用数变化未更新聚合/筛选标记的问题。\n"
          "Desktop list collapse and expansion remain on the left; selected details stay in the left pane.\n"
          "Sources: src/store/modules/stations.js, MapView.vue, mapService.js; map-markers.spec.js, viewport-debounce.spec.js.")

    # 07 — Realistic interaction order
    s = new_slide(prs, 7, "充电交互，遵循现实顺序", "A charging flow that makes sense.", "CHARGING STATE MACHINE")
    stages = [("明确选桩", "SELECT"), ("接枪确认", "PLUG IN"), ("检测启动", "CHECK & START"),
              ("停止应答", "STOP ACK"), ("拔枪归位", "UNPLUG"), ("查看账单", "CHECK OUT")]
    for i, (zh, en) in enumerate(stages):
        row, col = divmod(i, 3)
        x, y = 0.55 + col * 1.94, 2.34 + row * 1.34
        box(s, x, y, 1.72, 1.04, INK if i == 3 else WHITE, LINE if i != 3 else None)
        text(s, zh, x + 0.12, y + 0.19, 1.48, 0.33, size=17, fill=WHITE if i == 3 else INK, bold=True, align=PP_ALIGN.CENTER)
        text(s, en, x + 0.08, y + 0.62, 1.56, 0.22, size=10.5, fill=MINT if i == 3 else GRAY, font=EN, align=PP_ALIGN.CENTER)
        if col < 2:
            text(s, "→", x + 1.76, y + 0.32, 0.2, 0.3, size=15, fill=GREEN, font=EN)
    pair(s, "功率 × 时间，目标电量封顶", "kW × h → kWh, capped at the selected target", 0.6, 5.65, 5.72, size=18)
    phone(s, "mobile-plug.png", 6.83, 1.98, 2.23)
    phone(s, "mobile-unplug.png", 10.04, 1.98, 2.23)
    text(s, "先接枪，再开始 / Plug before start", 6.6, 6.83, 2.8, 0.22, size=10, fill=GRAY, align=PP_ALIGN.CENTER)
    text(s, "先停止，再拔枪 / Stop before unplug", 9.8, 6.83, 2.8, 0.22, size=10, fill=GRAY, align=PP_ALIGN.CENTER)
    notes(s, "只能从用户明确选择的充电桩开始。插枪确认后才执行虚拟设备连接检测与启动，并检查车辆接口、容量、功率和目标电量。\n"
          "达到目标或用户结束时请求停止；停止应答前仍保持 charging/stopping 保护，禁止拔枪或结算。停止后仍占用枪位，拔枪才释放并进入账单。\n"
          "本轮修复后台计时器延迟唤醒或停止应答等待导致的目标电量越界：模型只计选定目标内的有效输出电量与时间，仍保留实际经过时间。\n"
          "The model uses constant limited power for a demo; it does not simulate a production battery charging curve. All hardware acknowledgements are sandbox events.\n"
          "Sources: chargingModel.js, chargerGateway.js, charging store; charging-target.spec.js and charging-session.spec.js.")

    # 08 — Languages and layouts
    s = new_slide(prs, 8, "同一套体验，跨语言与屏幕", "One experience — across languages and screens.", "BILINGUAL & RESPONSIVE")
    photo(s, "english-detail-1440.png", 0.55, 2.04, 6.84)
    phone(s, "mobile-english.png", 8.0, 2.04, 2.05)
    phone(s, "mobile-charging-dark.png", 10.66, 2.04, 2.05)
    text(s, "桌面 / DESKTOP", 1.5, 6.85, 4.8, 0.21, size=10, fill=GRAY, font=EN)
    text(s, "中文 ⇄ English", 7.9, 6.63, 2.3, 0.28, size=12, fill=GREEN, align=PP_ALIGN.CENTER)
    text(s, "明暗一致 / LIGHT & DARK", 10.35, 6.63, 2.7, 0.28, size=11, fill=GRAY, align=PP_ALIGN.CENTER)
    notes(s, "中文和英文在渲染时切换，不改写订单中的标准站点、车辆或金额数据；金额仍为人民币。语言与主题保存到当前浏览器。\n"
          "本轮验证桌面 1440×1000、1024×768 与窄屏 390×844、320×568。九个账户页面复用统一结构，面板提供焦点管理、Esc 与忙碌锁。\n"
          "Desktop Edge viewport emulation is not a physical phone, Safari, screen-reader, or full WCAG certification.\n"
          "Sources: src/i18n, AccountPanel.vue, PanelShell.vue, scripts/capture-presentation.mjs, e2e/project.spec.js.")

    # 09 — Verification
    s = new_slide(prs, 9, "演示，也经得起验证", "Built to demo. Tested to hold up.", "QUALITY")
    metrics = [("427", "单元与组件测试", "Unit & component tests"), ("8", "浏览器回归", "Browser regression flows"),
               ("2,000", "站点契约校验", "Validated demo stations"), ("0", "依赖审计告警", "Known npm audit findings")]
    for i, (metric, zh, en) in enumerate(metrics):
        x = 0.55 + i * 3.14
        box(s, x, 2.12, 2.8, 2.72, INK if i == 0 else WHITE, None if i == 0 else LINE)
        text(s, metric, x + 0.23, 2.47, 2.4, 0.88, size=45, fill=NEON if i == 0 else GREEN, font=EN, bold=True)
        pair(s, zh, en, x + 0.23, 3.66, 2.33, size=17, fill=WHITE if i == 0 else INK, secondary=MINT if i == 0 else GRAY)
    text(s, "本轮修复 / FIXED IN THIS REVIEW", 0.58, 5.32, 11.9, 0.3, size=11, fill=GRAY, font=EN, bold=True)
    for x, zh, en in [(0.6, "目标电量封顶", "Target-limited metering"), (4.83, "预约兼容复核", "Booking compatibility checks"),
                      (9.0, "地图标记复用", "Correct marker lifecycle")]:
        pair(s, zh, en, x, 5.91, 3.8, size=19)
    notes(s, "验证日期：2026-10-28。45 个测试文件、427 项单元/组件测试通过；8 条 Playwright 浏览器流程通过。\n"
          "数据校验：20 个详情分片、2000 站、13 项契约断言。npm 官方源审计已安装依赖：337 个包、0 个已知漏洞。\n"
          "覆盖率仅统计 store/services：行 89.61%，分支 78.22%，函数 88.70%；不把该范围描述为全项目覆盖率。\n"
          "真实 SDK 的布局与主题验证另见 TEST_REPORT.md 和 live-review.json。测试通过不是不存在任何缺陷的保证。\n"
          "No production charger, real payment, physical GPS, or live driving-route acceptance is claimed. CI/deployment state must be checked separately from a local pass.")

    # 10 — Close
    s = new_slide(prs, 10, "", "", "TRY THE DEMO", dark=True)
    text(s, "可讲解。可演示。\n可扩展。", 0.446, 1.22, 7.4, 1.4, size=34, fill=NEON)
    text(s, "Explain it. Demo it. Build on it.", 0.446, 2.87, 11.8, 0.55, size=23, fill=MINT, font=EN)
    for x, zh, en in [(0.63, "有设计的数据", "DESIGNED DATA"), (4.87, "轻量双色地图", "ONE MAP, TWO THEMES"), (9.0, "连贯充电闭环", "CONNECTED JOURNEY")]:
        box(s, x, 3.64, 0.08, 0.62, NEON)
        pair(s, zh, en, x + 0.25, 3.57, 3.6, size=20, fill=WHITE, secondary=MINT)
    box(s, 0.63, 5.03, 4.05, 0.63, NEON)
    text(s, "打开演示 / OPEN LIVE DEMO  ↗", 0.84, 5.19, 3.64, 0.31, size=15, fill=INK, bold=True, link=DEMO)
    text(s, DEMO, 0.65, 5.96, 11.9, 0.37, size=21, fill=WHITE, font=EN, link=DEMO)
    text(s, "模拟站点与交易，不接入真实设备。 / Simulated stations & transactions. No live charger control.",
         0.65, 6.58, 12.0, 0.27, size=11, fill=MINT)
    notes(s, f"演示地址：{DEMO}\nRepository: https://github.com/mavorx/ChargingMap\n"
          "在线地址可能仍运行旧部署；本 PPT 的测试与截图来自本轮本地版本。Git 推送不等于 GitHub Pages 已部署。\n"
          "建议讲解顺序：先看站点与双色地图，再讲数据生成，最后实际演示选桩、插枪、停止、拔枪与支付。\n"
          "Future work, not current functionality: real operator protocols, payment providers, live availability and server-side accounts.")

    assert len(prs.slides) == 10
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (number, shape.name)
            assert shape.left + shape.width <= prs.slide_width + Inches(0.01), (number, shape.name)
            assert shape.top + shape.height <= prs.slide_height + Inches(0.01), (number, shape.name)
    validate_branding(prs, template)
    update_document_properties(prs)
    destination = OUT / "ChargingMap_Project_Overview_ZH-EN.pptx"
    prs.save(destination)
    print(f"Saved {destination} ({len(prs.slides)} slides)")
    return destination


def render(deck, soffice, preview_dir):
    import pymupdf

    approved_temp = Path(os.environ.get("LOCALAPPDATA", tempfile.gettempdir())) / "Temp" / "opencode"
    temporary_root = approved_temp if approved_temp.is_dir() else Path(tempfile.gettempdir())
    with tempfile.TemporaryDirectory(prefix="chargingmap-ppt-", dir=temporary_root) as temporary:
        profile = (Path(temporary) / "profile").as_uri()
        executable = Path(soffice)
        if os.name == "nt" and executable.with_suffix(".com").is_file():
            executable = executable.with_suffix(".com")
        subprocess.run([str(executable), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(OUT), str(deck)], check=True, timeout=150)
    pdf = deck.with_suffix(".pdf")
    assert pdf.stat().st_mtime_ns >= deck.stat().st_mtime_ns, "Renderer did not update the PDF"
    document = pymupdf.open(pdf)
    assert len(document) == 10, "PDF has missing slides"
    presentation = Presentation(deck)
    preview_dir = Path(preview_dir)
    preview_dir.mkdir(parents=True, exist_ok=True)
    thumbs = []
    for i, page in enumerate(document):
        assert len(page.get_text().strip()) > 60, f"Missing editable text on slide {i + 1}"
        for mark in logo_shapes(presentation.slides[i]):
            region = pymupdf.Rect(mark.left / 12700, mark.top / 12700,
                                 (mark.left + mark.width) / 12700, (mark.top + mark.height) / 12700)
            pixels = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=region, alpha=False)
            crop = Image.frombytes("RGB", (pixels.width, pixels.height), pixels.samples)
            dark_slide = i in {0, 2, 9}
            data = crop.tobytes()
            ink_pixels = sum(min(rgb) > 180 if dark_slide else max(rgb) < 140
                             for rgb in zip(data[::3], data[1::3], data[2::3]))
            assert ink_pixels > 80, f"Company mark did not render on PDF slide {i + 1}: {mark.name}"
        image = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
        filename = preview_dir / f"slide-{i + 1:02d}.png"
        image.save(filename)
        thumb = Image.open(filename).convert("RGB")
        thumb.thumbnail((800, 450), Image.Resampling.LANCZOS)
        thumbs.append(thumb.copy())
    montage = Image.new("RGB", (1648, 2348), "#DDE5E5")
    for i, thumb in enumerate(thumbs):
        montage.paste(thumb, (16 + (i % 2) * 816, 16 + (i // 2) * 466))
    montage.save(OUT / "ChargingMap_Preview.png", optimize=True)
    print(f"Rendered {pdf}; verified 20 visible mark regions; full-size previews: {preview_dir}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=TEMPLATE, help="User-provided VCTC POTX/PPTX template")
    parser.add_argument("--render-soffice", help="Optional LibreOffice soffice executable")
    parser.add_argument("--preview-dir", default=str(OUT / "preview"))
    arguments = parser.parse_args()
    result = build(arguments.template)
    if arguments.render_soffice:
        render(result, arguments.render_soffice, arguments.preview_dir)
